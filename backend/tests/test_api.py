"""Tests for FastAPI endpoints."""
import base64
import json
from unittest.mock import AsyncMock, MagicMock, patch

import pytest


def test_health_returns_200(test_client):
    resp = test_client.get("/api/health")
    assert resp.status_code == 200


def test_health_shape(test_client):
    resp = test_client.get("/api/health")
    data = resp.json()
    assert "status" in data
    # Ollama is unreachable in tests; chromadb is pre-set → status must be "degraded"
    assert data["status"] == "degraded"
    assert isinstance(data["ollama"], bool)
    assert isinstance(data["chromadb"], bool)
    assert isinstance(data["documents_count"], int)


def test_health_chromadb_ok_when_pre_set(test_client):
    """ChromaDB check should succeed because fixture injects in-memory client."""
    resp = test_client.get("/api/health")
    data = resp.json()
    assert data["chromadb"] is True
    assert data["documents_count"] == 0


def test_admin_documents_requires_auth(test_client):
    resp = test_client.get("/api/admin/documents")
    assert resp.status_code == 401


def test_admin_documents_with_auth(test_client, admin_headers):
    resp = test_client.get("/api/admin/documents", headers=admin_headers)
    assert resp.status_code == 200
    assert isinstance(resp.json(), list)


def test_admin_documents_wrong_password(test_client):
    creds = base64.b64encode(b"admin:wrongpassword").decode()
    headers = {"Authorization": f"Basic {creds}"}
    resp = test_client.get("/api/admin/documents", headers=headers)
    assert resp.status_code == 401


def test_admin_documents_wrong_user(test_client):
    creds = base64.b64encode(b"notadmin:changeme").decode()
    headers = {"Authorization": f"Basic {creds}"}
    resp = test_client.get("/api/admin/documents", headers=headers)
    assert resp.status_code == 401


def test_admin_documents_missing_auth_header(test_client):
    resp = test_client.get("/api/admin/documents", headers={})
    assert resp.status_code == 401


def test_admin_upload_requires_auth(test_client):
    files = {"file": ("test.txt", b"hello", "text/plain")}
    resp = test_client.post("/api/admin/upload", files=files)
    assert resp.status_code == 401


def test_admin_upload_rejects_txt(test_client, admin_headers):
    files = {"file": ("readme.txt", b"hello world", "text/plain")}
    resp = test_client.post("/api/admin/upload", files=files, headers=admin_headers)
    assert resp.status_code == 400


def test_admin_upload_rejects_jpg(test_client, admin_headers):
    files = {"file": ("photo.jpg", b"\xff\xd8\xff", "image/jpeg")}
    resp = test_client.post("/api/admin/upload", files=files, headers=admin_headers)
    assert resp.status_code == 400


def test_admin_upload_rejects_csv(test_client, admin_headers):
    files = {"file": ("data.csv", b"a,b,c\n1,2,3", "text/csv")}
    resp = test_client.post("/api/admin/upload", files=files, headers=admin_headers)
    assert resp.status_code == 400


def test_admin_upload_rejects_no_extension(test_client, admin_headers):
    files = {"file": ("noextension", b"data", "application/octet-stream")}
    resp = test_client.post("/api/admin/upload", files=files, headers=admin_headers)
    assert resp.status_code == 400


def test_admin_delete_requires_auth(test_client):
    resp = test_client.delete("/api/admin/documents/some-file-id")
    assert resp.status_code == 401


def test_admin_delete_nonexistent_file_id(test_client, admin_headers):
    """Deleting a non-existent file_id should return 404."""
    resp = test_client.delete(
        "/api/admin/documents/nonexistent-id", headers=admin_headers
    )
    assert resp.status_code == 404


def test_admin_reindex_requires_auth(test_client):
    resp = test_client.post("/api/admin/reindex")
    assert resp.status_code == 401


def test_admin_reindex_empty_dir(test_client, admin_headers):
    """Reindex with no documents returns zero counts."""
    resp = test_client.post("/api/admin/reindex", headers=admin_headers)
    assert resp.status_code == 200
    data = resp.json()
    assert data["files_processed"] == 0
    assert data["total_chunks"] == 0
    assert "duration_seconds" in data


def test_chat_503_when_no_chroma(test_client):
    """Chat endpoint returns 503 when ChromaDB is unavailable."""
    import main as main_module
    original = main_module.app.state.chroma_client
    try:
        del main_module.app.state.chroma_client
        resp = test_client.post("/api/chat", json={"question": "hello"})
        assert resp.status_code == 503
    finally:
        main_module.app.state.chroma_client = original


def test_admin_upload_rejects_oversized_file(test_client, admin_headers):
    """Upload endpoint returns 413 when file exceeds 50 MB."""
    large_content = b"x" * (50 * 1024 * 1024 + 1)
    files = {"file": ("large.pdf", large_content, "application/pdf")}
    resp = test_client.post("/api/admin/upload", files=files, headers=admin_headers)
    assert resp.status_code == 413


def test_admin_upload_returns_500_on_ingest_failure(test_client, admin_headers, tmp_path):
    """Upload endpoint returns 500 and cleans up file when ingestion fails."""
    import main as main_module

    minimal_pdf = b"%PDF-1.4\n1 0 obj\n<< /Type /Catalog >>\nendobj\n%%EOF"
    files = {"file": ("fail.pdf", minimal_pdf, "application/pdf")}

    with patch("main.ingest_file", side_effect=RuntimeError("embed failed")):
        resp = test_client.post("/api/admin/upload", files=files, headers=admin_headers)

    assert resp.status_code == 500
    # File should have been cleaned up
    dest = main_module.DOCUMENTS_DIR / "fail.pdf"
    assert not dest.exists()


def test_admin_upload_missing_filename(test_client, admin_headers):
    """Upload without a filename is rejected (400 from our check or 422 from FastAPI)."""
    files = {"file": ("", b"data", "application/pdf")}
    resp = test_client.post("/api/admin/upload", files=files, headers=admin_headers)
    assert resp.status_code in (400, 422)


def test_admin_upload_restores_old_file_on_failure(test_client, admin_headers):
    """When ingest fails for a re-upload, original file bytes are restored on disk."""
    import main as main_module

    original_content = b"original pdf bytes"
    dest = main_module.DOCUMENTS_DIR / "fail.pdf"
    dest.write_bytes(original_content)

    with patch("main.ingest_file", side_effect=RuntimeError("embed failed")):
        files = {"file": ("fail.pdf", b"new content", "application/pdf")}
        resp = test_client.post("/api/admin/upload", files=files, headers=admin_headers)

    assert resp.status_code == 500
    assert dest.exists()
    assert dest.read_bytes() == original_content


def test_admin_upload_returns_422_when_no_text_extracted(test_client, admin_headers):
    """Upload returns 422 when ingest succeeds but extracts zero chunks."""
    with patch("main.ingest_file", return_value={"filename": "empty.pdf", "chunks_created": 0, "file_id": "x"}):
        files = {"file": ("empty.pdf", b"%PDF-1.4", "application/pdf")}
        resp = test_client.post("/api/admin/upload", files=files, headers=admin_headers)
    assert resp.status_code == 422
    assert "No text" in resp.json()["detail"]


def test_admin_upload_rejects_filename_with_semicolon(test_client, admin_headers):
    """Upload rejects filenames containing shell-special characters like semicolons."""
    files = {"file": ("report;rm.pdf", b"%PDF-1.4", "application/pdf")}
    resp = test_client.post("/api/admin/upload", files=files, headers=admin_headers)
    assert resp.status_code == 400


def test_admin_upload_rejects_filename_with_special_chars(test_client, admin_headers):
    """Upload rejects filenames containing characters outside the allowed set."""
    files = {"file": ("report[final].pdf", b"%PDF-1.4", "application/pdf")}
    resp = test_client.post("/api/admin/upload", files=files, headers=admin_headers)
    assert resp.status_code == 400


def test_admin_delete_removes_file_from_disk(test_client, mock_chroma, admin_headers, dummy_embed, tmp_path):
    """Deleting a document via the API removes the file from disk."""
    import main as main_module
    from docx import Document as DocxDocument
    from rag.ingest import ingest_file

    doc_path = main_module.DOCUMENTS_DIR / "todelete.docx"
    doc = DocxDocument()
    doc.add_paragraph("This file should be removed on delete.")
    doc.save(str(doc_path))
    result = ingest_file(doc_path, mock_chroma, dummy_embed)

    resp = test_client.delete(
        f"/api/admin/documents/{result['file_id']}", headers=admin_headers
    )
    assert resp.status_code == 200
    assert not doc_path.exists(), "File was not removed from disk after delete"


def test_admin_list_503_when_no_chroma(test_client, admin_headers):
    """List endpoint returns 503 when ChromaDB is unavailable."""
    import main as main_module
    original = main_module.app.state.chroma_client
    try:
        del main_module.app.state.chroma_client
        resp = test_client.get("/api/admin/documents", headers=admin_headers)
        assert resp.status_code == 503
    finally:
        main_module.app.state.chroma_client = original


def test_admin_upload_503_when_no_chroma(test_client, admin_headers):
    """Upload endpoint returns 503 when ChromaDB is unavailable."""
    import main as main_module
    original = main_module.app.state.chroma_client
    try:
        del main_module.app.state.chroma_client
        files = {"file": ("test.pdf", b"%PDF-1.4", "application/pdf")}
        resp = test_client.post("/api/admin/upload", files=files, headers=admin_headers)
        assert resp.status_code == 503
    finally:
        main_module.app.state.chroma_client = original


def test_admin_delete_503_when_no_chroma(test_client, admin_headers):
    """Delete endpoint returns 503 when ChromaDB is unavailable."""
    import main as main_module
    original = main_module.app.state.chroma_client
    try:
        del main_module.app.state.chroma_client
        resp = test_client.delete("/api/admin/documents/some-id", headers=admin_headers)
        assert resp.status_code == 503
    finally:
        main_module.app.state.chroma_client = original


def test_admin_reindex_503_when_no_chroma(test_client, admin_headers):
    """Reindex endpoint returns 503 when ChromaDB is unavailable."""
    import main as main_module
    original = main_module.app.state.chroma_client
    try:
        del main_module.app.state.chroma_client
        resp = test_client.post("/api/admin/reindex", headers=admin_headers)
        assert resp.status_code == 503
    finally:
        main_module.app.state.chroma_client = original


def test_admin_reindex_partial_failure_returns_200(test_client, admin_headers):
    """Reindex with one success and one failure returns 200 with failed_files=1."""
    import main as main_module

    docs_dir = main_module.DOCUMENTS_DIR
    (docs_dir / "good.docx").write_bytes(b"placeholder")
    (docs_dir / "bad.docx").write_bytes(b"placeholder")

    def fake_ingest(path, *args, **kwargs):
        if path.name == "bad.docx":
            raise RuntimeError("ingest failed")
        return {"filename": path.name, "chunks_created": 3, "file_id": "abc123"}

    with patch("main.ingest_file", side_effect=fake_ingest):
        resp = test_client.post("/api/admin/reindex", headers=admin_headers)

    assert resp.status_code == 200
    data = resp.json()
    assert data["files_processed"] == 1
    assert data["failed_files"] == 1


def test_admin_reindex_all_fail_returns_500(test_client, admin_headers):
    """Reindex where every eligible file fails returns 500."""
    import main as main_module

    (main_module.DOCUMENTS_DIR / "fail.docx").write_bytes(b"placeholder")

    with patch("main.ingest_file", side_effect=RuntimeError("always fails")):
        resp = test_client.post("/api/admin/reindex", headers=admin_headers)

    assert resp.status_code == 500
    assert "all files failed" in resp.json()["detail"].lower()


def test_admin_reindex_zero_chunks_counted_as_failure(test_client, admin_headers):
    """Reindex counts zero-chunk results as failed files, not processed files."""
    import main as main_module

    docs_dir = main_module.DOCUMENTS_DIR
    (docs_dir / "good.pdf").write_bytes(b"placeholder")
    (docs_dir / "image_only.pdf").write_bytes(b"placeholder")

    def fake_ingest(path, *args, **kwargs):
        if path.name == "image_only.pdf":
            return {"filename": path.name, "chunks_created": 0, "file_id": "y"}
        return {"filename": path.name, "chunks_created": 5, "file_id": "x"}

    with patch("main.ingest_file", side_effect=fake_ingest):
        resp = test_client.post("/api/admin/reindex", headers=admin_headers)

    assert resp.status_code == 200
    data = resp.json()
    assert data["files_processed"] == 1
    assert data["failed_files"] == 1


def test_admin_documents_includes_disk_only_files(test_client, admin_headers):
    """Documents list includes files on disk that are not indexed in ChromaDB."""
    import main as main_module

    (main_module.DOCUMENTS_DIR / "unindexed.pdf").write_bytes(b"some pdf bytes")

    resp = test_client.get("/api/admin/documents", headers=admin_headers)
    assert resp.status_code == 200
    docs = resp.json()
    unindexed = [d for d in docs if d["filename"] == "unindexed.pdf"]
    assert len(unindexed) == 1
    assert unindexed[0]["file_id"] is None
    assert unindexed[0]["chunks"] == 0


def test_admin_reindex_clears_existing_collection(
    test_client, mock_chroma, admin_headers, dummy_embed, tmp_path
):
    """Reindex wipes existing collection entries even when no files are on disk."""
    from docx import Document as DocxDocument
    from rag.ingest import ingest_file, list_files

    doc_path = tmp_path / "dummy.docx"
    doc = DocxDocument()
    doc.add_paragraph("Content to be cleared on reindex.")
    doc.save(str(doc_path))
    ingest_file(doc_path, mock_chroma, dummy_embed)
    assert len(list_files(mock_chroma)) == 1

    # Remove the file from disk so reindex ingests nothing
    doc_path.unlink()

    resp = test_client.post("/api/admin/reindex", headers=admin_headers)
    assert resp.status_code == 200
    assert resp.json()["files_processed"] == 0
    assert list_files(mock_chroma) == []
